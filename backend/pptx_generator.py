"""Apollo PPTX generation adapted from the image-based export boundary used by NotebookLM-Lite.

Apollo currently renders source-grounded editable text slides directly with python-pptx.
The export boundary intentionally stays small so the generation workflow can evolve
toward image-backed slides later without changing the Studio API.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


APOLLO_BG = RGBColor(13, 14, 21)
APOLLO_TEXT = RGBColor(227, 225, 236)
APOLLO_MUTED = RGBColor(167, 139, 125)
APOLLO_ORANGE = RGBColor(249, 115, 22)
APOLLO_SOFT_ORANGE = RGBColor(255, 182, 144)


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font_size,
    color,
    bold=False,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def build_source_grounded_pptx(
    *,
    slides: list[dict[str, Any]],
    source_names: list[str],
    aspect_ratio: str = "16:9",
) -> bytes:
    """Create a real editable PPTX from structured source-grounded slide content.

    This follows the same small export boundary as NotebookLM-Lite's
    SlideDeckPPTXExporter: construct a Presentation, set its aspect ratio,
    add slides, and return the binary payload.
    """
    presentation = Presentation()
    if aspect_ratio == "4:3":
        presentation.slide_width = Inches(10)
        presentation.slide_height = Inches(7.5)
    else:
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)

    blank_layout = presentation.slide_layouts[6]

    for index, item in enumerate(slides):
        slide = presentation.slides.add_slide(blank_layout)
        _set_background(slide, APOLLO_BG)

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.62), Inches(0.9), Inches(0.08)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = APOLLO_ORANGE
        accent.line.fill.background()

        slide_title = str(item.get("title") or f"Slide {index + 1}").strip()
        _add_textbox(
            slide,
            Inches(0.55),
            Inches(0.9),
            Inches(12.0),
            Inches(0.75),
            slide_title,
            font_size=30 if index == 0 else 27,
            color=APOLLO_SOFT_ORANGE,
            bold=True,
        )

        bullets = [str(value).strip() for value in (item.get("bullets") or []) if str(value).strip()]
        if not bullets:
            bullets = ["No slide points were returned."]

        body = slide.shapes.add_textbox(
            Inches(0.72), Inches(1.85), Inches(11.9), Inches(4.65)
        )
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.04)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.03)
        frame.margin_bottom = Inches(0.03)

        for bullet_index, bullet in enumerate(bullets[:6]):
            paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
            paragraph.text = f"• {bullet}"
            paragraph.level = 0
            paragraph.font.size = Pt(19 if len(bullets) <= 4 else 16)
            paragraph.font.color.rgb = APOLLO_TEXT
            paragraph.space_after = Pt(11)

        notes = str(item.get("speaker_notes") or "").strip()
        if notes:
            notes_box = slide.shapes.add_textbox(
                Inches(0.72), Inches(6.42), Inches(11.9), Inches(0.42)
            )
            notes_frame = notes_box.text_frame
            notes_frame.clear()
            notes_run = notes_frame.paragraphs[0].add_run()
            notes_run.text = f"Notes: {notes}"
            notes_run.font.size = Pt(9)
            notes_run.font.color.rgb = APOLLO_MUTED

        source_footer = ", ".join(source_names[:4])
        if len(source_names) > 4:
            source_footer += f" + {len(source_names) - 4} more"

        _add_textbox(
            slide,
            Inches(0.72),
            Inches(7.03),
            Inches(9.2),
            Inches(0.22),
            f"Apollo · Source-grounded · {source_footer}" if source_footer else "Apollo · Source-grounded",
            font_size=8,
            color=APOLLO_MUTED,
        )
        _add_textbox(
            slide,
            Inches(11.3),
            Inches(7.03),
            Inches(1.2),
            Inches(0.22),
            str(index + 1),
            font_size=8,
            color=APOLLO_MUTED,
            align=PP_ALIGN.RIGHT,
        )

    if not presentation.slides:
        slide = presentation.slides.add_slide(blank_layout)
        _set_background(slide, APOLLO_BG)
        _add_textbox(
            slide,
            Inches(0.8),
            Inches(2.8),
            Inches(11.7),
            Inches(1),
            "Apollo Slide Deck",
            font_size=32,
            color=APOLLO_SOFT_ORANGE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
