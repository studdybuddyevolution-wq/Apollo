"""Apollo notebook/RAG service with token-aware chunks and hybrid retrieval."""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import shutil
import tempfile
import threading
import uuid
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader

from chunking import chunk_text, token_count
from embeddings import embed_text_sync
from storage import STORE

DATA_DIR = Path(os.getenv("APOLLO_DATA_DIR", Path(__file__).resolve().parent / "data"))
NOTEBOOKS_FILE = DATA_DIR / "notebooks.json"
_LOCK = threading.RLock()
_TOKEN_RE = re.compile(r"[A-Za-z0-9_]+")


def _now() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def _load_manifest() -> dict[str, list[dict[str, Any]]]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not NOTEBOOKS_FILE.exists():
        return {}
    try:
        return json.loads(NOTEBOOKS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _save_manifest(manifest: dict[str, list[dict[str, Any]]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    NOTEBOOKS_FILE.write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def _user_key(user_id: str | None) -> str:
    return (user_id or "default").strip() or "default"


def _notebook_dir(notebook_id: str) -> Path:
    return DATA_DIR / "notebooks" / notebook_id


def _metadata_path(notebook_id: str) -> Path:
    return _notebook_dir(notebook_id) / "chunks.json"

def _source_meta_path(notebook_id: str) -> Path:
    return _notebook_dir(notebook_id) / "sources.json"

def _source_payload_path(notebook_id: str, filename: str) -> Path:
    digest = hashlib.sha256(filename.encode("utf-8")).hexdigest()
    return _notebook_dir(notebook_id) / "payloads" / (digest + ".bin")

def sanitize_source_filename(filename: str) -> str:
    """Return a safe source name with no path components or control bytes."""
    raw = str(filename or "").replace("\\", "/").strip()
    safe = os.path.basename(raw)
    if not safe or safe in {".", ".."}:
        raise ValueError("Invalid source filename")
    if "\x00" in safe or any(ord(char) < 32 for char in safe):
        raise ValueError("Invalid source filename")
    return safe


def _source_candidate(filename: str, counter: int) -> str:
    if counter == 0:
        return filename
    path = Path(filename)
    return f"{path.stem} ({counter}){path.suffix}"


def _reserve_filesystem_source_name(notebook_id: str, requested_name: str) -> tuple[str, Path]:
    """Atomically claim a unique local source name for a development fallback."""
    notebook_dir = _notebook_dir(notebook_id)
    claims_dir = notebook_dir / ".claims"
    claims_dir.mkdir(parents=True, exist_ok=True)

    metadata = _load_source_meta(notebook_id)
    chunks = _load_chunks(notebook_id)
    used = set(metadata)
    used.update(str(item.get("source") or "") for item in chunks)

    counter = 0
    while True:
        candidate = _source_candidate(requested_name, counter)
        if candidate in used:
            counter += 1
            continue
        claim_path = claims_dir / (hashlib.sha256(candidate.encode("utf-8")).hexdigest() + ".claim")
        try:
            with claim_path.open("x", encoding="utf-8") as handle:
                handle.write(candidate)
            return candidate, claim_path
        except FileExistsError:
            counter += 1


def reserve_source_name(
    user_id: str | None,
    notebook_id: str,
    filename: str,
    *,
    kind: str = "file",
    source_url: str | None = None,
) -> tuple[str, Path | None]:
    """Reserve a unique source name before a new upload is processed."""
    if not get_notebook(user_id, notebook_id):
        raise KeyError("Notebook not found")
    requested = sanitize_source_filename(filename)
    if STORE:
        return STORE.reserve_source_name(notebook_id, requested, kind, _now(), source_url), None
    with _LOCK:
        return _reserve_filesystem_source_name(notebook_id, requested)


def _load_source_meta(notebook_id: str) -> dict[str, dict[str, Any]]:
    path = _source_meta_path(notebook_id)
    if not path.exists():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}

def _save_source_meta(notebook_id: str, meta: dict[str, dict[str, Any]]) -> None:
    notebook = _notebook_dir(notebook_id)
    notebook.mkdir(parents=True, exist_ok=True)
    _source_meta_path(notebook_id).write_text(json.dumps(meta, indent=2), encoding="utf-8")

def _save_source_payload_fs(notebook_id: str, filename: str, raw: bytes) -> None:
    path = _source_payload_path(notebook_id, filename)
    root = _notebook_dir(notebook_id).resolve()
    resolved = path.resolve()
    if not (resolved == root or root in resolved.parents):
        raise ValueError("Invalid source payload path")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=".payload-", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(raw)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise

def _load_source_payload_fs(notebook_id: str, filename: str) -> bytes | None:
    path = _source_payload_path(notebook_id, filename)
    return path.read_bytes() if path.exists() else None

def _delete_source_payload_fs(notebook_id: str, filename: str) -> None:
    _source_payload_path(notebook_id, filename).unlink(missing_ok=True)


def _tokens(text: str) -> list[str]:
    return [token.lower() for token in _TOKEN_RE.findall(text)]


def _clean_text(text: str) -> str:
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_text(text: str, chunk_size: int = 400, overlap: int = 60) -> list[str]:
    """Backward-compatible wrapper around the token-aware chunker."""
    return [chunk.text for chunk in chunk_text(text, chunk_tokens=chunk_size, overlap_tokens=overlap)]


def _extract_text(filename: str, raw: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if lower.endswith(".docx"):
        document = DocxDocument(io.BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return raw.decode("utf-8", errors="replace")


def _load_chunks(notebook_id: str) -> list[dict[str, Any]]:
    path = _metadata_path(notebook_id)
    if not path.exists():
        return []
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save_chunks(notebook_id: str, chunks: list[dict[str, Any]]) -> None:
    notebook = _notebook_dir(notebook_id)
    notebook.mkdir(parents=True, exist_ok=True)
    _metadata_path(notebook_id).write_text(json.dumps(chunks, indent=2), encoding="utf-8")


def list_notebooks(user_id: str | None = None) -> list[dict[str, Any]]:
    key = _user_key(user_id)
    if STORE:
        return STORE.list_notebooks(key)
    with _LOCK:
        return list(_load_manifest().get(key, []))


def create_notebook(user_id: str | None, title: str) -> dict[str, Any]:
    key = _user_key(user_id)
    title = title.strip() or "Untitled Notebook"
    now = _now()
    record = {"id": "nb_" + uuid.uuid4().hex[:10], "title": title, "created": now, "updated": now, "source_count": 0, "node_count": 0}
    if STORE:
        STORE.create_notebook(record, key)
        return record
    with _LOCK:
        manifest = _load_manifest()
        manifest.setdefault(key, []).append(record)
        _save_manifest(manifest)
    return record


def get_notebook(user_id: str | None, notebook_id: str) -> dict[str, Any] | None:
    key = _user_key(user_id)
    if STORE:
        return STORE.get_notebook(key, notebook_id)
    return next((nb for nb in list_notebooks(user_id) if nb["id"] == notebook_id), None)


def rename_notebook(user_id: str | None, notebook_id: str, title: str) -> dict[str, Any] | None:
    key = _user_key(user_id)
    clean_title = title.strip()
    if STORE:
        return STORE.rename_notebook(key, notebook_id, clean_title or "Untitled Notebook", _now())
    with _LOCK:
        manifest = _load_manifest()
        for notebook in manifest.get(key, []):
            if notebook["id"] == notebook_id:
                notebook["title"] = clean_title or notebook["title"]
                notebook["updated"] = _now()
                _save_manifest(manifest)
                return notebook
    return None


def delete_notebook(user_id: str | None, notebook_id: str) -> bool:
    key = _user_key(user_id)
    if STORE:
        return STORE.delete_notebook(key, notebook_id)
    with _LOCK:
        manifest = _load_manifest()
        notebooks = manifest.get(key, [])
        if not any(nb["id"] == notebook_id for nb in notebooks):
            return False
        manifest[key] = [nb for nb in notebooks if nb["id"] != notebook_id]
        _save_manifest(manifest)
        shutil.rmtree(_notebook_dir(notebook_id), ignore_errors=True)
        return True


def get_notebook_chunks(user_id: str | None, notebook_id: str, source_names: list[str] | None = None) -> list[dict[str, Any]]:
    """Load indexed chunks strictly from one notebook and optional sources."""
    if not get_notebook(user_id, notebook_id):
        return []
    chunks = STORE.list_chunks(notebook_id) if STORE else _load_chunks(notebook_id)
    allowed = {name for name in (source_names or []) if name}
    if allowed:
        chunks = [chunk for chunk in chunks if chunk.get("source") in allowed]
    return chunks


def list_sources(user_id: str | None, notebook_id: str) -> list[dict[str, Any]]:
    if not get_notebook(user_id, notebook_id):
        return []
    if STORE:
        rows = STORE.list_sources_with_status(notebook_id)
        if rows:
            return rows
        return STORE.list_sources(notebook_id)
    chunks = _load_chunks(notebook_id)
    metadata = _load_source_meta(notebook_id)
    grouped: dict[str, dict[str, Any]] = {}
    for chunk in chunks:
        source = chunk["source"]
        meta = metadata.get(source, {})
        grouped.setdefault(
            source,
            {
                "name": source,
                "kind": meta.get("kind") or chunk.get("kind", "file"),
                "chunks": 0,
                "status": meta.get("status", "indexed"),
                "error": meta.get("error"),
                "source_url": meta.get("source_url"),
            },
        )
        grouped[source]["chunks"] += 1
    for source, meta in metadata.items():
        grouped.setdefault(
            source,
            {
                "name": source,
                "kind": meta.get("kind", "file"),
                "chunks": 0,
                "status": meta.get("status", "pending"),
                "error": meta.get("error"),
                "source_url": meta.get("source_url"),
            },
        )
    return sorted(grouped.values(), key=lambda item: item["name"].lower())


def get_source_payload(user_id: str | None, notebook_id: str, filename: str) -> bytes | None:
    if not get_notebook(user_id, notebook_id):
        return None
    if STORE:
        try:
            return STORE.get_source_payload(notebook_id, filename)
        except Exception:
            return None
    return _load_source_payload_fs(notebook_id, filename)


def get_source_metadata(user_id: str | None, notebook_id: str, filename: str) -> dict[str, Any] | None:
    if not get_notebook(user_id, notebook_id):
        return None
    if STORE:
        try:
            return STORE.get_source_metadata(notebook_id, filename)
        except Exception:
            return None
    return _load_source_meta(notebook_id).get(filename)


def add_source(user_id: str | None, notebook_id: str, filename: str, raw: bytes, *, kind: str = "file", source_url: str | None = None, replace_existing: bool = True) -> dict[str, Any]:
    if not get_notebook(user_id, notebook_id):
        raise KeyError("Notebook not found")

    filename = sanitize_source_filename(filename)
    claim_path: Path | None = None
    if not replace_existing:
        filename, claim_path = reserve_source_name(user_id, notebook_id, filename, kind=kind, source_url=source_url)

    now = _now()
    if STORE:
        STORE.upsert_source_status(notebook_id, filename, kind, "processing", now=now, source_url=source_url)
        try:
            STORE.save_source_payload(notebook_id, filename, raw, now)
            text = _extract_text(filename, raw)
            tokenized = chunk_text(text, filename=filename)
            if not tokenized:
                raise ValueError("No readable text found in source")
            new_chunks = [
                {"id": uuid.uuid4().hex, "source": filename, "kind": kind, "text": item.text, "chunk_index": item.index, "content_type": item.content_type}
                for item in tokenized
            ]
            STORE.replace_source(notebook_id, filename, new_chunks)
            STORE.update_counts(notebook_id, now, len(STORE.list_sources(notebook_id)), len(STORE.list_chunks(notebook_id)))
            STORE.upsert_source_status(notebook_id, filename, kind, "indexed", now=now, source_url=source_url)
            return {"name": filename, "kind": kind, "chunks": len(new_chunks), "characters": len(text), "tokens": token_count(text), "status": "indexed", "source_url": source_url}
        except Exception as exc:
            STORE.upsert_source_status(notebook_id, filename, kind, "failed", str(exc)[:500], now=_now(), source_url=source_url)
            raise
        finally:
            if claim_path:
                claim_path.unlink(missing_ok=True)

    try:
        with _LOCK:
            metadata = _load_source_meta(notebook_id)
            metadata[filename] = {"kind": kind, "status": "processing", "error": None, "source_url": source_url, "updated": now}
            _save_source_meta(notebook_id, metadata)
            _save_source_payload_fs(notebook_id, filename, raw)

        try:
            text = _extract_text(filename, raw)
            tokenized = chunk_text(text, filename=filename)
            if not tokenized:
                raise ValueError("No readable text found in source")
            new_chunks = [
                {"id": uuid.uuid4().hex, "source": filename, "kind": kind, "text": item.text, "chunk_index": item.index, "content_type": item.content_type}
                for item in tokenized
            ]
        except Exception as exc:
            with _LOCK:
                metadata = _load_source_meta(notebook_id)
                metadata[filename] = {"kind": kind, "status": "failed", "error": str(exc)[:500], "source_url": source_url, "updated": _now()}
                _save_source_meta(notebook_id, metadata)
            raise

        with _LOCK:
            existing = [c for c in _load_chunks(notebook_id) if c.get("source") != filename]
            all_chunks = existing + new_chunks
            _save_chunks(notebook_id, all_chunks)
            metadata = _load_source_meta(notebook_id)
            metadata[filename] = {"kind": kind, "status": "indexed", "error": None, "source_url": source_url, "updated": now}
            _save_source_meta(notebook_id, metadata)
            key = _user_key(user_id)
            manifest = _load_manifest()
            for notebook in manifest.get(key, []):
                if notebook["id"] == notebook_id:
                    notebook["updated"] = now
                    notebook["source_count"] = len(list_sources(user_id, notebook_id))
                    notebook["node_count"] = len(all_chunks)
                    break
            _save_manifest(manifest)
        return {"name": filename, "kind": kind, "chunks": len(new_chunks), "characters": len(text), "tokens": token_count(text), "status": "indexed", "source_url": source_url}
    finally:
        if claim_path:
            claim_path.unlink(missing_ok=True)


def remove_source(user_id: str | None, notebook_id: str, filename: str) -> bool:
    if not get_notebook(user_id, notebook_id):
        return False
    if STORE:
        removed = STORE.delete_source(notebook_id, filename)
        if removed:
            STORE.update_counts(notebook_id, _now(), len(STORE.list_sources(notebook_id)), len(STORE.list_chunks(notebook_id)))
        return removed
    with _LOCK:
        existing = _load_chunks(notebook_id)
        chunks = [c for c in existing if c.get("source") != filename]
        metadata = _load_source_meta(notebook_id)
        existed = len(chunks) != len(existing) or filename in metadata or _load_source_payload_fs(notebook_id, filename) is not None
        if not existed:
            return False
        _save_chunks(notebook_id, chunks)
        metadata.pop(filename, None)
        _save_source_meta(notebook_id, metadata)
        _delete_source_payload_fs(notebook_id, filename)
    return True


def _bm25_scores(query: str, chunks: list[dict[str, Any]]) -> list[tuple[float, int]]:
    if not chunks:
        return []
    query_terms = _tokens(query)
    if not query_terms:
        return []
    document_terms = [_tokens(chunk.get("text", "")) for chunk in chunks]
    doc_lengths = [len(tokens) for tokens in document_terms]
    avgdl = sum(doc_lengths) / max(1, len(doc_lengths))
    document_frequency: Counter[str] = Counter()
    for terms in document_terms:
        document_frequency.update(set(terms))
    scores: list[tuple[float, int]] = []
    k1 = 1.5
    b = 0.75
    n_docs = len(chunks)
    for index, terms in enumerate(document_terms):
        term_counts = Counter(terms)
        dl = doc_lengths[index]
        score = 0.0
        for term in query_terms:
            df = document_frequency.get(term, 0)
            if not df:
                continue
            tf = term_counts.get(term, 0)
            if not tf:
                continue
            idf = math.log(1 + (n_docs - df + 0.5) / (df + 0.5))
            denom = tf + k1 * (1 - b + b * dl / max(avgdl, 1.0))
            score += idf * ((tf * (k1 + 1)) / max(denom, 1e-9))
        scores.append((score, index))
    scores.sort(reverse=True)
    return scores


def retrieve_hybrid(
    user_id: str | None,
    notebook_id: str,
    query: str,
    top_k: int = 5,
    source_names: list[str] | None = None,
    rrf_k: int = 60,
) -> list[dict[str, Any]]:
    """Strictly notebook-scoped BM25 + vector retrieval with reciprocal-rank fusion."""
    if not get_notebook(user_id, notebook_id):
        return []
    chunks = get_notebook_chunks(user_id, notebook_id, source_names)
    if not chunks:
        return []

    candidate_k = max(top_k * 3, 10)
    bm25_ranked = _bm25_scores(query, chunks)[:candidate_k]
    by_id = {chunk.get("id"): chunk for chunk in chunks}
    fused: dict[str, dict[str, Any]] = {}
    for rank, (score, index) in enumerate(bm25_ranked, 1):
        chunk = chunks[index]
        key = str(chunk.get("id") or f"bm25-{index}")
        fused[key] = {
            "id": chunk.get("id"),
            "source": chunk.get("source", "unknown source"),
            "text": chunk.get("text", ""),
            "score": 1.0 / (rrf_k + rank),
            "bm25_score": float(score),
            "vector_score": 0.0,
            "rrf_score": 1.0 / (rrf_k + rank),
            "retrieval_method": "bm25",
        }

    vector_ranked: list[dict[str, Any]] = []
    if STORE and STORE.vector_available() and os.getenv("GEMINI_API_KEY", "").strip():
        try:
            query_embedding = embed_text_sync(query)
            vector_ranked = STORE.vector_search(notebook_id, query_embedding, candidate_k, source_names)
        except Exception as exc:
            print(f"[Apollo retrieval] vector search unavailable; using BM25: {exc}")

    for rank, item in enumerate(vector_ranked, 1):
        key = str(item.get("id") or f"vector-{rank}")
        current = fused.get(key)
        if current is None:
            chunk = by_id.get(item.get("id"))
            current = {
                "id": item.get("id"),
                "source": item.get("source", chunk.get("source") if chunk else "unknown source"),
                "text": item.get("text", chunk.get("text", "") if chunk else ""),
                "score": 0.0,
                "bm25_score": 0.0,
                "vector_score": 0.0,
                "rrf_score": 0.0,
                "retrieval_method": "vector",
            }
            fused[key] = current
        current["vector_score"] = float(item.get("vector_score", 0.0))
        current["rrf_score"] += 1.0 / (rrf_k + rank)
        current["score"] = current["rrf_score"]
        current["retrieval_method"] = "hybrid" if current["bm25_score"] > 0 else "vector"

    results = sorted(fused.values(), key=lambda item: (item["rrf_score"], item["vector_score"], item["bm25_score"]), reverse=True)
    return results[:top_k]


def retrieve(user_id: str | None, notebook_id: str, query: str, top_k: int = 5, source_names: list[str] | None = None) -> list[dict[str, Any]]:
    return retrieve_hybrid(user_id, notebook_id, query, top_k=top_k, source_names=source_names)


def format_context(results: list[dict[str, Any]], max_chars: int = 9000, max_tokens: int | None = None) -> str:
    if not results:
        return ""
    blocks: list[str] = []
    used_chars = 0
    used_tokens = 0
    for idx, result in enumerate(results, 1):
        block = f"[Source {idx}: {result['source']}]\n{result['text']}"
        block_tokens = token_count(block)
        if max_tokens is not None and blocks and used_tokens + block_tokens > max_tokens:
            break
        if used_chars + len(block) > max_chars and blocks and max_tokens is None:
            break
        blocks.append(block)
        used_chars += len(block)
        used_tokens += block_tokens
    return "\n\n".join(blocks)
