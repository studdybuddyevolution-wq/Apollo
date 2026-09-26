"""Study Report v2 generation and rendering helpers for Marklyf.

The module keeps report planning, normalization, deterministic Markdown rendering,
and DOCX export separate from the HTTP route layer. It deliberately reuses
Marklyf's existing notebook RAG and hybrid research stack rather than introducing
a second research framework.
"""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from diagrams import content_overlap_ratio
from phase3_common import extract_json_object
from research_engine import (
    TOPIC_TEMPLATES,
    build_outline,
    classify_query,
    extract_requested_sections,
    format_evidence,
    run_hybrid_research,
)


ACADEMIC_SECTIONS = [
    "Abstract",
    "Introduction & Background",
    "Literature / Context",
    "Methodology",
    "Findings & Analysis",
    "Discussion & Limitations",
    "Conclusion & Future Work",
]

REPORT_MODES = {"study", "deep", "academic"}
MAX_SECTIONS = 8


def normalize_sections(values: Iterable[Any], *, fallback: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = re.sub(r"\s+", " ", str(value or "")).strip(" .:;")
        if len(text) < 3:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(text[:140])
        if len(result) >= MAX_SECTIONS:
            break
    return result or fallback[:MAX_SECTIONS]


def build_report_plan(
    focus: str,
    mode: str,
    requested_sections: list[str] | None = None,
) -> list[str]:
    clean_focus = (focus or "the selected notebook sources").strip()
    explicit = normalize_sections(requested_sections or [], fallback=[])
    if explicit:
        return explicit
    if mode == "academic":
        return ACADEMIC_SECTIONS.copy()
    requested_from_question = extract_requested_sections(clean_focus)
    if requested_from_question:
        return normalize_sections(requested_from_question, fallback=TOPIC_TEMPLATES["conceptual"])
    topic = classify_query(clean_focus)
    return normalize_sections(
        build_outline(topic, clean_focus),
        fallback=TOPIC_TEMPLATES["conceptual"],
    )


def _clean_source_refs(values: Any, allowed_sources: list[str]) -> list[str]:
    allowed = {str(source).strip(): str(source).strip() for source in allowed_sources}
    refs: list[str] = []
    if isinstance(values, str):
        values = [values]
    for value in values or []:
        ref = str(value or "").strip()
        if ref in allowed and ref not in refs:
            refs.append(ref)
    return refs


def normalize_report_data(
    data: dict[str, Any],
    *,
    plan: list[str],
    allowed_sources: list[str],
) -> dict[str, Any]:
    sections: list[dict[str, Any]] = []
    raw_sections = data.get("sections") if isinstance(data, dict) else []
    if isinstance(raw_sections, list):
        for index, raw in enumerate(raw_sections[:MAX_SECTIONS], 1):
            if not isinstance(raw, dict):
                continue
            heading = str(raw.get("heading") or "").strip() or (plan[index - 1] if index <= len(plan) else f"Section {index}")
            content = str(raw.get("content") or "").strip()
            points = [
                str(point).strip()
                for point in (raw.get("points") or [])
                if str(point).strip()
            ][:8]
            sections.append(
                {
                    "heading": heading[:160],
                    "content": content,
                    "points": points,
                    "source_refs": _clean_source_refs(raw.get("source_refs"), allowed_sources),
                }
            )
    if not sections:
        sections = [
            {
                "heading": heading,
                "content": "",
                "points": [],
                "source_refs": [],
            }
            for heading in plan
        ]

    key_terms: list[dict[str, Any]] = []
    for raw in (data.get("key_terms") or []) if isinstance(data, dict) else []:
        if not isinstance(raw, dict):
            continue
        term = str(raw.get("term") or "").strip()
        definition = str(raw.get("definition") or "").strip()
        if term and definition:
            key_terms.append(
                {
                    "term": term[:100],
                    "definition": definition,
                    "source_refs": _clean_source_refs(raw.get("source_refs"), allowed_sources),
                }
            )
        if len(key_terms) >= 12:
            break

    takeaways = [
        str(item).strip()
        for item in (data.get("exam_takeaways") or []) if str(item).strip()
    ][:10]

    questions: list[dict[str, str]] = []
    for raw in (data.get("exam_questions") or []) if isinstance(data, dict) else []:
        if isinstance(raw, dict):
            question = str(raw.get("question") or "").strip()
            answer = str(raw.get("answer") or "").strip()
            if question and answer:
                questions.append({"question": question, "answer": answer})
        elif str(raw).strip():
            questions.append({"question": str(raw).strip(), "answer": ""})
        if len(questions) >= 10:
            break

    return {
        "title": str(data.get("title") or "Marklyf Study Report").strip() or "Marklyf Study Report",
        "summary": str(data.get("summary") or "").strip(),
        "sections": sections,
        "key_terms": key_terms,
        "exam_takeaways": takeaways,
        "exam_questions": questions,
    }


def render_report_markdown(
    data: dict[str, Any],
    *,
    source_names: list[str],
    web_sources: list[dict[str, str]] | None = None,
) -> str:
    lines = [
        f"# {data.get('title') or 'Marklyf Study Report'}",
        "",
        str(data.get("summary") or "").strip(),
    ]

    sections = [section for section in (data.get("sections") or []) if isinstance(section, dict)]
    if sections:
        lines.extend(["", "## Table of Contents", ""])
        for index, section in enumerate(sections, 1):
            heading = str(section.get("heading") or f"Section {index}").strip()
            lines.append(f"{index}. {heading}")

    for section in sections:
        heading = str(section.get("heading") or "Section").strip()
        lines.extend(["", f"## {heading}"])
        content = str(section.get("content") or "").strip()
        if content:
            lines.extend(["", content])
        for point in section.get("points") or []:
            value = str(point).strip()
            if value:
                lines.append(f"- {value}")
        refs = _clean_source_refs(section.get("source_refs"), source_names)
        if refs:
            lines.extend(["", "Sources: " + ", ".join(refs)])

    key_terms = data.get("key_terms") or []
    if key_terms:
        lines.extend(["", "## Key Terms", ""])
        for item in key_terms:
            if isinstance(item, dict):
                term = str(item.get("term") or "").strip()
                definition = str(item.get("definition") or "").strip()
                if term and definition:
                    lines.append(f"- **{term}:** {definition}")

    takeaways = [str(item).strip() for item in (data.get("exam_takeaways") or []) if str(item).strip()]
    if takeaways:
        lines.extend(["", "## Exam-Oriented Takeaways", ""])
        lines.extend(f"- {item}" for item in takeaways)

    questions = data.get("exam_questions") or []
    if questions:
        lines.extend(["", "## Practice Questions", ""])
        for item in questions:
            if isinstance(item, dict):
                lines.append(f"### {item.get('question', '').strip()}")
                answer = str(item.get("answer") or "").strip()
                if answer:
                    lines.extend(["", f"**Answer:** {answer}"])

    web_sources = web_sources or []
    if web_sources:
        lines.extend(["", "## References", ""])
        seen_urls: set[str] = set()
        for index, source in enumerate(web_sources, 1):
            url = str(source.get("url") or "").strip()
            title = str(source.get("title") or url or f"Web source {index}").strip()
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
            lines.append(f"{len(seen_urls)}. [{title}]({url})")
    elif source_names:
        lines.extend(["", "## References", ""])
        for source in source_names:
            lines.append(f"- Notebook source: {source}")

    return "\n".join(lines).strip() + "\n"


def report_grounding_score(markdown: str, context: str) -> float:
    return round(content_overlap_ratio("text", markdown, context), 2)


def build_full_report_prompt(
    *,
    focus: str,
    mode: str,
    plan: list[str],
    source_names: list[str],
    context: str,
    evidence: str = "",
) -> tuple[str, str, int]:
    outline = "\n".join(f"{index}. {heading}" for index, heading in enumerate(plan, 1))
    allowed = ", ".join(source_names) or "the selected notebook sources"
    mode_instruction = {
        "study": "Optimize for revision: clear explanations, definitions, examples, distinctions, key terms, and exam-oriented takeaways.",
        "academic": (
            "Use a formal academic structure. Never invent methodology, results, statistics, experiments, citations, or findings. "
            "If the supplied sources do not support a section such as Methodology or Findings, explicitly say so rather than fabricating content."
        ),
        "deep": (
            "Use the supplied notebook + web evidence. Prefer primary and authoritative evidence, surface meaningful conflicts, "
            "and cite claims with the numbered evidence references."
        ),
    }[mode]
    evidence_block = f"\n\nNUMBERED RESEARCH EVIDENCE:\n{evidence}" if evidence else ""
    system = (
        "You are Marklyf's Study Report engine. Produce a polished, source-grounded report. "
        "Never invent facts. Keep every source_refs value restricted to the allowed notebook source names. "
        "Return ONLY valid JSON."
    )
    prompt = f"""
Create a {mode} report for this focus:
{focus}

REPORT OUTLINE:
{outline}

{mode_instruction}

Return JSON with exactly these top-level keys:
title, summary, sections, key_terms, exam_takeaways, exam_questions

sections: array of objects with heading, content, points, source_refs.
- Use one section for each planned heading, in order.
- content should be 2-4 substantive paragraphs.
- points should be concise only when useful.
- source_refs must contain only allowed notebook source names.
- In deep mode, inline citations such as [1] must refer to the numbered evidence blocks supplied below.

key_terms: 5-12 objects with term, definition, source_refs.
exam_takeaways: 5-10 concise bullets.
exam_questions: 3-8 objects with question and answer.

ALLOWED NOTEBOOK SOURCES:
{allowed}

NOTEBOOK SOURCE CONTEXT:
{context}
{evidence_block}
"""
    return prompt.strip(), system, 3600 if mode != "deep" else 3200


def build_section_prompt(
    *,
    focus: str,
    mode: str,
    section_heading: str,
    source_names: list[str],
    context: str,
    evidence: str = "",
) -> tuple[str, str, int]:
    allowed = ", ".join(source_names) or "the selected notebook sources"
    extra = (
        "Use only notebook evidence."
        if mode != "deep"
        else "Use notebook + numbered web evidence. Cite important web-supported claims with [n]."
    )
    prompt = f"""
Write ONE section for a Marklyf {mode} Study Report.

Focus: {focus}
Section: {section_heading}

{extra}
Do not invent facts, examples, dates, methodology, results, statistics, or citations.
Return ONLY JSON:
{{
  "heading": "...",
  "content": "2-4 substantive paragraphs",
  "points": ["optional concise point", "..."],
  "source_refs": ["allowed notebook source name", "..."],
  "key_terms": [{{"term":"...","definition":"...","source_refs":[]}}],
  "exam_takeaways": ["..."]
}}

ALLOWED NOTEBOOK SOURCES:
{allowed}

NOTEBOOK SOURCE CONTEXT:
{context}
"""
    if evidence:
        prompt += f"\n\nNUMBERED RESEARCH EVIDENCE:\n{evidence}"
    return prompt.strip(), (
        "You are Marklyf's section-level Study Report writer. Stay strictly grounded in supplied evidence."
    ), 1700


def build_questions_prompt(focus: str, report_sections: list[dict[str, Any]]) -> tuple[str, str, int]:
    section_text = "\n\n".join(
        f"## {item.get('heading', '')}\n{item.get('content', '')}\n" for item in report_sections
    )
    prompt = f"""
Create 4-8 exam-oriented questions for this study report.

Focus: {focus}

Return ONLY JSON:
{{"exam_questions":[{{"question":"...","answer":"..."}}],"summary":"..."}}.

Questions must be answerable from the supplied report sections. Do not introduce outside facts.

REPORT:
{section_text}
"""
    return prompt.strip(), "You create source-grounded study questions from supplied report content.", 1200


def normalize_section_response(data: dict[str, Any], heading: str, source_names: list[str]) -> dict[str, Any]:
    normalized = normalize_report_data(
        {
            "title": "Section",
            "summary": "",
            "sections": [data],
            "key_terms": data.get("key_terms") if isinstance(data, dict) else [],
            "exam_takeaways": data.get("exam_takeaways") if isinstance(data, dict) else [],
            "exam_questions": [],
        },
        plan=[heading],
        allowed_sources=source_names,
    )
    return normalized["sections"][0]


def report_to_docx(markdown: str, title: str = "Marklyf Study Report") -> bytes:
    document = Document()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title or "Marklyf Study Report")
    run.bold = True
    run.font.size = Pt(20)

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if not line:
            document.add_paragraph()
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        else:
            document.add_paragraph(re.sub(r"\*\*(.*?)\*\*", r"\1", line))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


__all__ = [
    "REPORT_MODES",
    "ACADEMIC_SECTIONS",
    "build_report_plan",
    "build_full_report_prompt",
    "build_section_prompt",
    "build_questions_prompt",
    "normalize_report_data",
    "normalize_section_response",
    "render_report_markdown",
    "report_grounding_score",
    "report_to_docx",
    "format_evidence",
    "run_hybrid_research",
]
